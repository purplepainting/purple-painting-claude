import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
m=json.load(open('work/measure.json'))
wb=Workbook()
F=lambda **k: Font(name='Arial',size=10,**k)
H=PatternFill('solid',fgColor='2C1654'); HF=Font(name='Arial',size=10,bold=True,color='FFFFFF')
fills={'Contract':'D6F0DD','Base – assumed':'FFF4CC','Alt / RFI':'FDE3CC','Not in contract':'E5E5E5'}
thin=Side(style='thin',color='CCCCCC'); B=Border(bottom=thin)
def header(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(row,i,c); x.font=HF; x.fill=H; x.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[chr(64+i)].width=w
# ---------------- BoQ ----------------
ws=wb.active; ws.title='BoQ'
ws['A1']='Weinberg Garage – Painting Bill of Quantities'; ws['A1'].font=Font(name='Arial',size=14,bold=True)
ws['A2']='2934 Quail Valley Road, Solvang CA | Arch 5/13/26, Struct 5/12/26 | Prepared by Purple Painting 10/15/2026 | Quantities measured from CAD linework; see Measurements tab'
ws['A3']='Enter unit rates in column H (blue, yellow fill). Status: Contract = called out to paint; Base – assumed = implied, price in base and confirm; Alt / RFI = price as add/alternate until answered; Not in contract = excluded.'
for c in ('A2','A3'): ws[c].font=F(italic=True)
cols=['#','Section','Description','Paint system (proposed – not specified)','Unit','Qty','Status','Rate ($/unit)','Amount ($)','Drawing ref','RFI','Note']
header(ws,5,cols,[5,12,44,34,7,9,15,11,12,22,6,46])
T=m['tot']; E=m['ext']
gyp_gross=T['gyp_lo']; openings=303
rows=[
('Interior','Gyp walls up to 12\'','Prime + 2 coats eggshell (Level 5 noted)','SF',f'={gyp_gross:.0f}-{openings}','Base – assumed','KN-1.0 FN-gwb-01; A-2.1','1','Gross 1,395 SF less 303 SF openings (3 OH doors 10x8, 2 doors 3x8, 2 windows 2.5x3). Eave 10\'-0" per plate/curb; measured at 9\'-10" – add ~2%.'),
('Interior','Gyp walls 12\' to 14\'-10" (gables, ceiling step faces)','Prime + 2 coats eggshell','SF',f'={T["gyp_hi"]:.0f}+{m["step_sf"]:.0f}','Base – assumed','A-5.1 sections; A-2.2','1','High work – lift. Gable triangles integrated + vertical step faces between offset bays.'),
('Interior','Vaulted gyp ceiling 4:12','Prime + 2 coats flat','SF',f'={sum(v["clg_sf"] for v in m["rooms"].values()):.0f}','Base – assumed','A-5.1; A-2.2','2','Floor 1,360 SF x 1.054. If trusses with flat ceiling: ~1,360 SF and gable walls drop out.'),
('Interior','Skylight wells (Velux VS A06)','Prime + 2 coats flat','EA','=3','Base – assumed','A-4.1 window sched A/B/C','','~13 SF gyp each, not detailed.'),
('Interior','Interior window casing, casement D & E','Prime + 2 coats semi-gloss','LF','=18','Base – assumed','A-4.13/14','4',''),
('Interior','Interior face 12" CMU walls (e4)','Block filler + 2 coats','SF',f'={T["cmu_lo"]+T["cmu_hi"]:.0f}','Alt / RFI','A-2.1 wall schedule e4','3','Finish not stated: paint, clear seal, or natural.'),
('Doors','Custom wood segmental garage doors 10\'x8\', both faces','Prime + 2 coats (or stain)','EA','=3','Alt / RFI','A-4.1 doors 2,3,4','4','Finish columns blank.'),
('Doors','Lincoln clad glazed doors 3\'x8\', interior face + casing','Prime + 2 coats semi-gloss (or stain)','EA','=2','Alt / RFI','A-4.1 doors 1,5','4','Clad exterior is factory finish.'),
('Exterior','2x wood fascia, painted','Prime + 2 coats exterior satin','LF',f'={E["fascia"]:.0f}','Contract','KN-1.0 WD-tr-01','7',f'Eave {E["eave"]:.0f} LF + exposed rake ~{E["rake"]:.0f} LF. Structural says 2x8 min.'),
('Exterior','26 ga galv. flashing at fascia, painted','Metal primer + 2 coats exterior','LF',f'={E["fascia"]:.0f}','Contract','A-4.11 rake eave detail','7','Color to match fascia or roof – confirm.'),
('Exterior','4x6 false rafter tails','Prime + 2 coats (or stain)','EA',f'={E["tails16"]:.0f}','Alt / RFI','KN-1.0 WD-rr-01; S-sheet eave','5',f'16" o.c. per arch = {E["tails16"]:.0f}; 30" o.c. per structural = {E["tails30"]:.0f}.'),
('Exterior','Overhang T&G soffit + exposed rake rafters','Prime + 2 coats (or stain)','SF','=530','Alt / RFI','A-4.11; S-sheet rake','6','Roof-plan overhang outline, sloped (Hakim 528). Claude first calc 676 ran high.'),
('Exterior','Post and rail fence','TBD','LF','=0','Alt / RFI','KN-1.0 EXT-fn-01','9','Not measured. Finish not stated.'),
('Flatwork','Garage slab sealer','Penetrating sealer','SF','=1360','Alt / RFI','KN-1.0 CC-slb-01','10','Not called out.'),
('Flatwork','Exterior slab sealer','Penetrating sealer','SF','=1857','Alt / RFI','KN-1.0 CC-slb-02','10','Not measured by Claude; Hakim qty.'),
('Exterior','Exterior CMU strip above retaining','Block filler + 2 coats','SF','=101','Alt / RFI','A-2.1 / A-5.1','3','Not measured by Claude; Hakim qty.'),
('Excluded','Hardie board & batt siding, batts','—','—','=0','Not in contract','KN-1.0 TMP-rsp-01/02','8','Integral color, pre-finished.'),
('Excluded','Gutters','—','—','=0','Not in contract','KN-1.0 TMP-rf-01',' ','Integral color.'),
('Excluded','Stone veneer, concrete retaining walls','—','—','=0','Not in contract','KN-1.0 MA-st-01',' ',''),
('Excluded','Exposed galvanized Simpson hardware','—','—','=0','Not in contract','S0.2 note 41','11','Leave unpainted unless directed.'),
]
r0=6
for i,(sec,desc,sysm,unit,qty,st,ref,rfi,note) in enumerate(rows):
    r=r0+i
    vals=[i+1,sec,desc,sysm,unit,qty,st,None,f'=IF(H{r}="","",F{r}*H{r})',ref,rfi,note]
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
    ws.cell(r,7).fill=PatternFill('solid',fgColor=fills[st])
    h=ws.cell(r,8); h.fill=PatternFill('solid',fgColor='FFFF00'); h.font=Font(name='Arial',size=10,color='0000FF'); h.number_format='$#,##0.00'
    ws.cell(r,9).number_format='$#,##0;($#,##0);-'; ws.cell(r,6).number_format='#,##0'
rl=r0+len(rows)-1
r=rl+2
ws.cell(r,3,'Totals by status').font=F(bold=True)
for k,st in enumerate(['Contract','Base – assumed','Alt / RFI']):
    ws.cell(r+1+k,3,st).font=F(); ws.cell(r+1+k,3).fill=PatternFill('solid',fgColor=fills[st])
    c=ws.cell(r+1+k,9,f'=SUMIFS(I{r0}:I{rl},G{r0}:G{rl},C{r+1+k})'); c.number_format='$#,##0;($#,##0);-'; c.font=F()
ws.cell(r+4,3,'Base bid (Contract + Base – assumed)').font=F(bold=True)
c=ws.cell(r+4,9,f'=I{r+1}+I{r+2}'); c.number_format='$#,##0;($#,##0);-'; c.font=F(bold=True)
ws.freeze_panes='A6'
ws.cell(6,6).comment=Comment('Openings 303 SF = 3 x (10x8) + 2 x (3x8) + 2 x (2.5x3). Source: A-4.1 schedules.','Claude')
# ---------------- Measurements ----------------
ms=wb.create_sheet('Measurements')
header(ms,1,['Room','Wall from (ft x,y)','Wall to (ft x,y)','Type','Share','LF','SF up to 12\'','SF above 12\''],[8,18,18,8,7,8,12,12])
r=2
for k,v in m['rooms'].items():
    for e in v['edges']:
        for typ,f in e.get('split',[[e['type'],1]]):
            for j,val in enumerate([k,str(e['a']),str(e['b']),typ.upper(),f,round(e['lf']*f,2),round(e['lo']*f,1),round(e['hi']*f,1)],1):
                ms.cell(r,j,val).font=F()
            r+=1
ms.cell(r,1,'Totals').font=F(bold=True)
for col in 'FGH': ms[f'{col}{r}']=f'=SUM({col}2:{col}{r-1})'; ms[f'{col}{r}'].font=F(bold=True)
ms.cell(r+2,1,'Ceiling step faces between bays (SF)').font=F(); ms.cell(r+2,7,round(m['step_sf'],1)).font=F()
ms.cell(r+3,1,'Ceiling sloped SF: 101 / 102').font=F(); ms.cell(r+3,7,m['rooms']['101']['clg_sf']).font=F(); ms.cell(r+3,8,m['rooms']['102']['clg_sf']).font=F()
ms.cell(r+5,1,'Method: walls traced from A-2.1 linework (1px = 1in at 1/4" scale), face-by-face height integration along 4:12 gable profile, eave 9\'-10" (bid uses 10\'-0"). Gross of openings.').font=F(italic=True)
# ---------------- Spec Review ----------------
sp=wb.create_sheet('Spec Review')
header(sp,1,['Source','Keynote / note','Text (as written)','Status','Observation','RFI'],[16,14,52,15,52,6])
spec=[
('KN-1.0','FN-gwb-01','(N) 5/8" WR Gyp. Bd., Level 5 Finish, typical at wet conditions. Prep as required for paint finish u.n.o.','Base – assumed','Only line implying paint. No paint spec, colors, sheens or system. Level 5 = high-build primer.','1'),
('KN-1.0','WD-tr-01','(N) 2x Wood Facia, painted.','Contract','Only explicit "painted" keynote.','7'),
('A-4.11','Rake eave detail','26 Ga. Gal. Flashing, painted.','Contract','Flashing painted; color not given.','7'),
('KN-1.0','WD-rr-01','(N) 4x6 false rafter tails at 16" o.c. max. See Structural.','Alt / RFI','Structural eave detail says 30" o.c. max. Finish not stated.','5'),
('A-4.11','Eave detail','1x6 T&G Sheathing @ Overhang','Alt / RFI','Structural shows 2x6 T&G starter board and exposed 2x6 T&G rake rafters. Finish not stated.','6'),
('KN-1.0','TMP-rsp-01/02','Hardi-Plank Fiber-Cement sheathing, w/ vertical batts, integral color; color 1','Not in contract','Pre-finished. Confirm touch-up by installer.','8'),
('KN-1.0','TMP-rf-01','5" Bevel gutter... integral color','Not in contract','',''),
('KN-1.0','MA-st-01','(N) Stone veneer on Concrete wall.','Not in contract','',''),
('A-2.1','Wall schedule e4','(N) EXTERIOR WALL: 12" CMU Wall with #6 bars at 8" o.c.','Alt / RFI','No interior finish listed for e4 (e1 lists 5/8" gyp). Interior CMU faces exposed.','3'),
('A-4.1','Door schedule','Door finish / frame finish columns blank','Alt / RFI','Custom wood segmental OH doors and Lincoln clad doors: finish and responsibility unknown.','4'),
('A-5.1','Sections','FN-gwb-01 at roof slope (vaulted)','Alt / RFI','Structural S-sheets: LAYOUT TRUSS, GIRDER TRUSS TO CMU. Vault vs flat ceiling changes ceiling SF and gable walls.','2'),
('A-5.1','Level markers','FF 360\'-0" / Top of Curb 357\'-4 11/16" / Top Plate 367\'-4 11/16"','Alt / RFI','Plate is 10\'-0" above curb but 7\'-4" above FF datum. Bid at 10\'-0".','12'),
('S0.2','Note 41','GALVANIZING: ALL EXPOSED STEEL TIMBER HARDWARE','Not in contract','Leave galvanized unless told.','11'),
('S0.2','Steel note 5','STRUCTURAL STEEL WORK DOES NOT REQUIRE SHOP PAINTING','Not in contract','No exposed steel found in garage.',''),
('KN-1.0','CC-slb-01/02','4" Concrete Garage slab / 4" Concrete slab','Alt / RFI','No sealer called.','10'),
('KN-1.0','EXT-fn-01','(N) Post and Rail Fence with 4" x 4" wire mesh','Alt / RFI','Finish not stated.','9'),
]
for i,row in enumerate(spec,2):
    for j,v in enumerate(row,1):
        c=sp.cell(i,j,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
    sp.cell(i,4).fill=PatternFill('solid',fgColor=fills[row[3]])
# ---------------- Materials ----------------
mt=wb.create_sheet('Materials')
mt['A1']='Specified materials (from drawings) and proposed coating systems. No paint spec in the documents – coating systems are Purple Painting proposals for pricing, to be submitted for approval.'; mt['A1'].font=F(italic=True)
header(mt,2,['Surface','Specified material (as drawn)','Source','Proposed primer','Proposed finish','Sheen','Color','Status'],[22,38,18,30,30,12,12,15])
mats=[
('Interior gyp walls','5/8" WR gypsum board, Level 5','KN-1.0 FN-gwb-01','High-build drywall primer-surfacer (Level 5)','2 coats interior acrylic','Eggshell','TBD (RFI 1)','Base – assumed'),
('Vaulted gyp ceiling','5/8" gyp on rafters/trusses','A-5.1','PVA / drywall primer','2 coats interior ceiling','Flat','TBD','Base – assumed'),
('Interior CMU (e4)','12" CMU, solid grout','A-2.1 e4','Block filler','2 coats interior acrylic or clear sealer','Eggshell','TBD','Alt / RFI'),
('Wood fascia','2x wood fascia (2x8 #1 per struct)','WD-tr-01, S-sheets','Exterior oil/stain-blocking wood primer, back-prime','2 coats exterior acrylic','Satin','TBD','Contract'),
('Metal flashing','26 ga galvanized flashing','A-4.11','Galvanized metal primer (after wash/etch)','2 coats exterior acrylic','Satin','Match fascia?','Contract'),
('Rafter tails, T&G soffit, rake rafters','4x6 DF tails; 1x6 or 2x6 T&G','KN-1.0, A-4.11, S-sheets','Exterior wood primer or stain','2 coats paint, or semi-transparent stain','Satin','TBD','Alt / RFI'),
('Garage doors','Custom wood-framed segmental doors 10x8','A-4.1','Exterior wood primer (all 6 faces/edges)','2 coats exterior, or stain + clear','Satin','TBD','Alt / RFI'),
('Man doors / windows','Lincoln clad glazed doors, clad casements','A-4.1','Interior wood primer','2 coats interior enamel or stain','Semi-gloss','TBD','Alt / RFI'),
('Concrete slab','4" concrete slab','CC-slb-01/02','—','Penetrating silane/siloxane sealer','—','Clear','Alt / RFI'),
('Siding','Hardie board & batt, integral color','TMP-rsp-01/02','—','By siding installer','—','Color 1','Not in contract'),
('Gutters','5" bevel gutter, integral color','TMP-rf-01','—','—','—','—','Not in contract'),
('Skylights','Velux VS A06 self-curbing, bronze','A-4.1','—','Well: gyp per ceiling system','—','—','Base – assumed'),
]
for i,row in enumerate(mats,3):
    for j,v in enumerate(row,1):
        c=mt.cell(i,j,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
    mt.cell(i,8).fill=PatternFill('solid',fgColor=fills[row[7]])
# ---------------- RFIs ----------------
rq=wb.create_sheet('RFIs')
header(rq,1,['#','Subject','Question','Reference','Answer','Date answered'],[5,26,70,26,30,12])
rfis=[
(1,'Paint scope and colors','No finish schedule or paint spec. Confirm surfaces to be painted, colors and sheens for interior gyp walls, ceiling, trim, fascia and flashing. Level 5 noted on FN-gwb-01; confirm primer-surfacer acceptable.','KN-1.0 FN-gwb-01, WD-tr-01'),
(2,'Vaulted ceiling vs trusses','A-5.1 shows vaulted gyp ceiling following 4:12 roof; structural shows prefab trusses and girder truss to CMU. Confirm scissor/vaulted trusses with gyp at slope, or flat ceiling at plate.','A-5.1; S-sheet roof framing'),
(3,'Interior face of CMU','Wall type e4 lists no interior finish. Furred and boarded, painted with block filler, clear sealed, or natural?','A-2.1 wall schedule e4'),
(4,'Door and window finishes','Finish columns blank. Confirm finish and responsibility for custom wood segmental garage doors (both faces), interior faces of Lincoln clad doors 1 & 5, casement windows D & E.','A-4.1'),
(5,'Rafter tails','Arch WD-rr-01 16" o.c.; structural 30" o.c. max. Confirm spacing and finish.','KN-1.0; S-sheet eave detail'),
(6,'Overhang sheathing, rake rafters','A-4.11 1x6 T&G; structural 2x6 T&G starter board and exposed 2x6 T&G rake rafters. Confirm material and finish.','A-4.11; S-sheet rake detail'),
(7,'Fascia size and color','Keynote 2x fascia painted; structural 2x8 #1 min. Confirm size, color; flashing color.','WD-tr-01; A-4.11'),
(8,'Hardie siding','Integral color board-and-batt: confirm factory pre-finished, touch-up by installer, no field paint.','TMP-rsp-01/02'),
(9,'Fence','Post and rail fence – any finish?','EXT-fn-01'),
(10,'Concrete sealer','No sealer called on garage or exterior slab. Confirm none, or specify.','CC-slb-01/02'),
(11,'Exposed hardware','Galvanized timber hardware: leave unpainted or paint to match?','S0.2 note 41'),
(12,'Floor datum','FF 360\'-0" vs curb 357\'-4 11/16" vs plate 367\'-4 11/16". Confirm slab elevation for wall heights.','A-5.1'),
]
for i,row in enumerate(rfis,2):
    for j,v in enumerate(row,1):
        c=rq.cell(i,j,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
    for j in (5,6): rq.cell(i,j).fill=PatternFill('solid',fgColor='FFFF00')
wb.save('Weinberg Garage - Painting Takeoff - Claude - 2026-10-15.xlsx')
print('ok')
